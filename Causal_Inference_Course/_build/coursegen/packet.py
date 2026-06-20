"""Build a weekly self-study packet (.docx) from a week-content dict.

The layout mirrors the Week 1 packet that ships with the course: a banner,
"this week in one sentence", learning goals, guided reading, a concept
refresher, a problem set with solutions, a lab, and a self-check / look-ahead.
"""
from __future__ import annotations

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from . import theme

# ---------------------------------------------------------------------------
# low-level helpers
# ---------------------------------------------------------------------------

def _rgb(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr)


def _shade(element, hexstr: str) -> None:
    """Apply a solid background fill to a paragraph or table cell."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexstr)
    element.append(shd)


def _set_cell_bg(cell, hexstr: str) -> None:
    _shade(cell._tc.get_or_add_tcPr(), hexstr)


def _para_bg(p, hexstr: str) -> None:
    _shade(p._p.get_or_add_pPr(), hexstr)


def _no_space(p) -> None:
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)


def _add_border(p, color: str, size: int = 18, where=("left",)) -> None:
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    for side in where:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "8")
        el.set(qn("w:color"), color)
        pbdr.append(el)
    pPr.append(pbdr)


# ---------------------------------------------------------------------------
# styled paragraph builders
# ---------------------------------------------------------------------------

def _runs(p, segments, size=11, color=theme.INK, bold=False, italic=False,
          font=theme.FONT_BODY):
    """segments: str or list of (text, overrides-dict)."""
    if isinstance(segments, str):
        segments = [(segments, {})]
    for text, ov in segments:
        r = p.add_run(text)
        r.font.name = ov.get("font", font)
        r.font.size = Pt(ov.get("size", size))
        r.font.bold = ov.get("bold", bold)
        r.font.italic = ov.get("italic", italic)
        r.font.color.rgb = _rgb(ov.get("color", color))
    return p


def banner(doc, week_no, title):
    p = doc.add_paragraph()
    _no_space(p)
    _runs(p, f"{theme.COURSE_TITLE}  ·  Self-study packet",
          size=11, color=theme.BLUE, bold=True)
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(2)
    h.paragraph_format.space_after = Pt(2)
    _runs(h, f"Week {week_no} — {title}", size=24, color=theme.NAVY,
          bold=True, font=theme.FONT_HEAD)
    # accent rule
    rule = doc.add_paragraph()
    _no_space(rule)
    _add_border(rule, theme.AMBER, size=24, where=("bottom",))


def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    _runs(p, text, size=16, color=theme.NAVY, bold=True, font=theme.FONT_HEAD)
    _add_border(p, theme.BLUE, size=18, where=("left",))
    p.paragraph_format.left_indent = Inches(0.08)
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    _runs(p, text, size=13, color=theme.BLUE, bold=True, font=theme.FONT_HEAD)
    return p


def body(doc, segments, italic=False, space=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space)
    _runs(p, segments, size=11, color=theme.INK, italic=italic)
    return p


def bullet(doc, segments, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    _runs(p, segments, size=11, color=theme.INK)
    return p


def numbered(doc, n, title_seg, body_seg=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    _runs(p, [(f"{n}. ", {"bold": True, "color": theme.BLUE})])
    _runs(p, title_seg if isinstance(title_seg, list) else [(title_seg, {"bold": True})])
    if body_seg:
        b = doc.add_paragraph()
        b.paragraph_format.space_after = Pt(4)
        _runs(b, body_seg)
    return p


def callout(doc, title, lines, color=theme.TEAL):
    """A shaded panel with a coloured left border — used for key ideas."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(0)
    _para_bg(p, theme.GREY_LIGHT)
    _add_border(p, color, size=24, where=("left",))
    _runs(p, title, size=11, color=color, bold=True)
    for ln in lines:
        q = doc.add_paragraph()
        _no_space(q)
        q.paragraph_format.space_after = Pt(2)
        _para_bg(q, theme.GREY_LIGHT)
        _add_border(q, color, size=24, where=("left",))
        _runs(q, ln, size=10.5, color=theme.INK)
    spacer = doc.add_paragraph()
    _no_space(spacer)


def code_block(doc, caption, code):
    if caption:
        c = doc.add_paragraph()
        c.paragraph_format.space_before = Pt(6)
        c.paragraph_format.space_after = Pt(0)
        _runs(c, caption, size=9.5, color=theme.GREY, bold=True)
    for line in code.rstrip("\n").split("\n"):
        p = doc.add_paragraph()
        _no_space(p)
        _para_bg(p, "F4F6F8")
        r = p.add_run(line if line else " ")
        r.font.name = theme.FONT_MONO
        r.font.size = Pt(9.5)
        r.font.color.rgb = _rgb("1B3A57")
    spacer = doc.add_paragraph()
    _no_space(spacer)
    spacer.paragraph_format.space_after = Pt(4)


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for j, htext in enumerate(headers):
        cell = t.rows[0].cells[j]
        _set_cell_bg(cell, theme.NAVY)
        cell.paragraphs[0].clear()
        _runs(cell.paragraphs[0], htext, size=10.5, color="FFFFFF", bold=True)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for j, val in enumerate(row):
            if ri % 2 == 1:
                _set_cell_bg(cells[j], theme.GREY_LIGHT)
            cells[j].paragraphs[0].clear()
            _runs(cells[j].paragraphs[0], str(val), size=10.5, color=theme.INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


# ---------------------------------------------------------------------------
# document margins / base style
# ---------------------------------------------------------------------------

def _setup(doc):
    style = doc.styles["Normal"]
    style.font.name = theme.FONT_BODY
    style.font.size = Pt(11)
    style.font.color.rgb = _rgb(theme.INK)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)


# ---------------------------------------------------------------------------
# main entry point
# ---------------------------------------------------------------------------

def build_packet(week: dict, out_path: str) -> str:
    doc = Document()
    _setup(doc)

    banner(doc, week["number"], week["title"])
    body(doc, week["packet_intro"], italic=True)

    h2(doc, "This week in one sentence")
    body(doc, week["one_sentence"])

    h1(doc, week["objectives_heading"])
    for item in week["objectives"]:
        bullet(doc, item)

    # --- guided reading -----------------------------------------------------
    h1(doc, "Guided reading")
    body(doc, week["reading_intro"], italic=True)
    h2(doc, "Core (do these)")
    for r in week["readings"]:
        bullet(doc, [(r["text"] + "  ", {"bold": True}),
                     ("What to look for: " + r["look_for"], {})])
    if week.get("optional_readings"):
        h2(doc, "Optional (pick one)")
        for r in week["optional_readings"]:
            bullet(doc, [(r["text"] + "  ", {"bold": True}), (r["note"], {})])

    # --- concept refresher --------------------------------------------------
    h1(doc, "Concept refresher")
    body(doc, week["concept_intro"], italic=True)
    for sec in week["concept_sections"]:
        h2(doc, sec["heading"])
        if sec.get("body"):
            body(doc, sec["body"])
        for b in sec.get("bullets", []):
            bullet(doc, b)
        if sec.get("callout"):
            callout(doc, sec["callout"]["title"], sec["callout"]["lines"],
                    color=getattr(theme, sec["callout"].get("color", "TEAL")))

    # --- problem set --------------------------------------------------------
    ps = week["problem_set"]
    h1(doc, f"{ps['label']} — {ps['title']}")
    body(doc, ps["intro"], italic=True)
    for i, prob in enumerate(ps["problems"], 1):
        numbered(doc, i, prob["title"], prob["prompt"])

    h1(doc, f"{ps['label']} — Solutions")
    body(doc, "Don't read until you've attempted every problem.", italic=True)
    for i, prob in enumerate(ps["problems"], 1):
        numbered(doc, i, prob["solution_title"])
        for ln in prob["solution"]:
            bullet(doc, ln)

    # --- lab ----------------------------------------------------------------
    lab = week["lab"]
    h1(doc, f"{lab['label']} — {lab['title']}")
    body(doc, [("Goal:  ", {"bold": True}), (lab["goal"], {})])
    for step in lab["steps"]:
        h2(doc, step["heading"])
        if step.get("body"):
            body(doc, step["body"])
        for b in step.get("bullets", []):
            bullet(doc, b)
        if step.get("code_r"):
            code_block(doc, "R", step["code_r"])
        if step.get("code_python"):
            code_block(doc, "Python", step["code_python"])
    if lab.get("expected"):
        h2(doc, "What you should see")
        body(doc, lab["expected"])
    if lab.get("submit"):
        h2(doc, "Submit")
        for s in lab["submit"]:
            bullet(doc, s)

    # --- self-check ---------------------------------------------------------
    h1(doc, "Self-check & looking ahead")
    h2(doc, f"You're ready for Week {week['number'] + 1 if week['number'] < 15 else 'the capstone'} if you can…")
    for s in week["self_check"]:
        bullet(doc, s)
    nxt = week["next_week"]
    h2(doc, nxt["heading"])
    body(doc, nxt["teaser"])
    body(doc, week.get("stuck_note",
         "Stuck? Post in the course forum or bring it to office hours — "
         "confusion now is normal and worth surfacing early."), italic=True)

    doc.save(out_path)
    return out_path

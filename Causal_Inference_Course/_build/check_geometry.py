#!/usr/bin/env python3
"""Geometry QA for built decks: flag any shape that extends past the slide
bounds (off-canvas) — a proxy for visual overflow, since this environment can't
render slides to images. Read-only; operates on the built .pptx files.
"""
from __future__ import annotations

import glob
import os
import sys

from pptx import Presentation
from pptx.util import Emu

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TOL = Emu(int(0.06 * 914400))   # ~0.06" tolerance for rounding


def check_deck(path: str) -> list[str]:
    prs = Presentation(path)
    W, H = prs.slide_width, prs.slide_height
    out = []
    for i, slide in enumerate(prs.slides):
        for sh in slide.shapes:
            try:
                l, t, w, h = sh.left, sh.top, sh.width, sh.height
            except Exception:
                continue
            if l is None or t is None or w is None or h is None:
                continue
            name = (sh.name or "shape")
            if l < -TOL or t < -TOL:
                out.append(f"slide {i+1}: '{name}' off top/left "
                           f"(l={l/914400:.2f}in t={t/914400:.2f}in)")
            if l + w > W + TOL:
                out.append(f"slide {i+1}: '{name}' past right edge "
                           f"(right={(l+w)/914400:.2f}in > {W/914400:.2f}in)")
            if t + h > H + TOL:
                out.append(f"slide {i+1}: '{name}' past bottom edge "
                           f"(bottom={(t+h)/914400:.2f}in > {H/914400:.2f}in)")
    return out


def main(weeks):
    dirs = sorted(glob.glob(os.path.join(ROOT, "Week*")))
    if weeks:
        dirs = [d for d in dirs if any(
            os.path.basename(d).startswith(f"Week{w:02d}_") for w in weeks)]
    total = 0
    for d in dirs:
        for f in glob.glob(os.path.join(d, "*_Lecture.pptx")):
            ws = check_deck(f)
            if ws:
                print(f"\n{os.path.basename(d)}:")
                for w in ws:
                    print("  ⚠ " + w)
            total += len(ws)
    print(f"\n{total} geometry warning(s).")
    return 0


if __name__ == "__main__":
    sys.exit(main([int(a) for a in sys.argv[1:]]))

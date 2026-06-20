#!/usr/bin/env python3
"""Validate generated artifacts: open the .docx and .pptx, and EXECUTE every
code cell of each .ipynb in a fresh namespace so we know the notebooks actually
run. Exit non-zero on any failure.
"""
from __future__ import annotations

import glob
import io
import json
import os
import sys
import traceback
import contextlib

import matplotlib
matplotlib.use("Agg")  # headless

from docx import Document
from pptx import Presentation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def check_docx(path: str) -> str:
    doc = Document(path)
    n = len(doc.paragraphs)
    if n < 40:
        raise AssertionError(f"{path}: only {n} paragraphs")
    return f"docx OK ({n} paragraphs, {len(doc.tables)} tables)"


def check_pptx(path: str) -> str:
    prs = Presentation(path)
    n = len(prs.slides)
    if n < 20:
        raise AssertionError(f"{path}: only {n} slides")
    return f"pptx OK ({n} slides)"


def run_ipynb(path: str) -> str:
    nb = json.load(open(path))
    code = [("".join(c["source"]) if isinstance(c["source"], list)
             else c["source"])
            for c in nb["cells"] if c["cell_type"] == "code"]
    ns: dict = {}
    buf = io.StringIO()
    for i, src in enumerate(code):
        try:
            with contextlib.redirect_stdout(buf):
                exec(compile(src, f"{os.path.basename(path)}#cell{i}", "exec"), ns)
        except Exception:
            raise AssertionError(
                f"{path}: code cell {i} raised:\n{traceback.format_exc()}")
    return f"ipynb OK ({len(code)} code cells executed)"


def main(weeks):
    failures = []
    dirs = sorted(glob.glob(os.path.join(ROOT, "Week*")))
    if weeks:
        dirs = [d for d in dirs if any(
            os.path.basename(d).startswith(f"Week{w:02d}_") for w in weeks)]
    for d in dirs:
        name = os.path.basename(d)
        print(f"\n=== {name} ===")
        for pat, fn in ((".docx", check_docx), (".pptx", check_pptx),
                        (".ipynb", run_ipynb)):
            for f in glob.glob(os.path.join(d, f"*{pat}")):
                try:
                    print(f"  {fn(f)}")
                except Exception as e:
                    print(f"  ✗ FAIL {os.path.basename(f)}: {e}")
                    failures.append(f)
    print("\n" + ("ALL GOOD" if not failures
                  else f"{len(failures)} FAILURE(S)"))
    return 1 if failures else 0


if __name__ == "__main__":
    wk = [int(a) for a in sys.argv[1:]]
    sys.exit(main(wk))
